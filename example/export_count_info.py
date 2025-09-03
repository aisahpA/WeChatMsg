import os
import openpyxl


def add_count_info_to_excel(count_info):
    """
    将统计信息添加到Excel中
    :param count_info: 统计信息
    :return: None
    """

    
    filename = os.path.join(count_info['output_dir'], "聊天记录导出统计信息.xlsx")

    if os.path.exists(filename):
        workbook = openpyxl.load_workbook(filename)
        sheet = workbook['聊天记录统计']
    else:
        workbook = openpyxl.Workbook()
        sheet = workbook.create_sheet("聊天记录统计", 0)
        columns = ['wxid', '联系人', '开始时间', '结束时间', 
                    '消息数量', '总文件量', '总存储', '总存储bytes',
                    '图片应该量', '图片实际量', '图片存储', '图片存储bytes',
                    '视频应该量', '视频实际量', '视频存储', '视频存储bytes',
                    '文件应该量', '文件实际量', '文件存储', '文件存储bytes',
                    '音频应该量', '音频实际量', '音频存储', '音频存储bytes'
                ]
        sheet.append(columns)
    
    # 统计文件数量和存储空间
    origin_path = count_info['origin_path']
    datas = []
    
    def append_by_type(should_count, path):
        middle = count_files_and_size(path)
        datas.append(should_count)
        datas.append(middle['file_count'])
        datas.append(middle['total_size_formatted'])
        datas.append(middle['total_size_bytes'])

    datas.append(count_info['wxid'])
    datas.append(count_info['contact_remark'])
    datas.append(count_info['start_time_str'])
    datas.append(count_info['end_time_str'])
    append_by_type(count_info['message_count'], origin_path)
    append_by_type(count_info['image_count'], os.path.join(origin_path, 'image'))
    append_by_type(count_info['video_count'], os.path.join(origin_path, 'video'))
    append_by_type(count_info['file_count'], os.path.join(origin_path, 'file'))
    append_by_type(count_info['audio_count'], os.path.join(origin_path, 'voice'))

    sheet.append(datas)   
    workbook.save(filename)


def count_files_and_size(folder_path):
    """
    统计指定文件夹下的文件数量和存储大小
    
    Args:
        folder_path (str): 文件夹路径
    
    Returns:
        dict: 包含文件数量和总大小的字典
    """
    file_count = 0
    total_size = 0
    
    # 遍历文件夹中的所有文件和子文件夹
    for root, dirs, files in os.walk(folder_path):
        file_count += len(files)
        for file in files:
            file_path = os.path.join(root, file)
            # 累加文件大小
            try:
                total_size += os.path.getsize(file_path)
            except OSError:
                # 处理无法访问的文件
                pass
    
    # 将字节转换为更易读的单位
    def format_size(size_bytes):
        if size_bytes == 0:
            return "0 B"
        size_names = ["B", "KB", "MB", "GB", "TB"]
        i = 0
        while size_bytes >= 1024 and i < len(size_names) - 1:
            size_bytes /= 1024.0
            i += 1
        return f"{size_bytes:.2f} {size_names[i]}"
    
    return {
        'file_count': file_count,
        'total_size_bytes': total_size,
        'total_size_formatted': format_size(total_size)
    }


